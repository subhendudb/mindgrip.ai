#!/usr/bin/env python3
"""Build a short MindGrip / SO-101 project deck from repo docs (max 10 slides)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "MindGrip_SO101_Project.pptx"
BRAND = ROOT / "docs-web" / "public" / "brand" / "mindgrip-logo.png"
HERO = ROOT / "docs-web" / "public" / "images" / "home-sorting-scene.png"
SCENE = ROOT / "docs-web" / "public" / "images" / "so101-camera-placement.png"
CONTROL = ROOT / "docs-web" / "public" / "images" / "s3_control_flow.png"

# Industrial slate + teal (not purple / cream AI defaults)
INK = RGBColor(0x14, 0x1C, 0x24)
SLATE = RGBColor(0x2A, 0x3A, 0x48)
MUTED = RGBColor(0x5A, 0x6B, 0x7A)
TEAL = RGBColor(0x0D, 0x8A, 0x7A)
TEAL_DEEP = RGBColor(0x0A, 0x5C, 0x52)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF2, 0xF5, 0xF7)
ACCENT_LINE = RGBColor(0x0D, 0x8A, 0x7A)


def _set_run(run, *, size=18, bold=False, color=INK, font="Calibri"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _add_textbox(slide, left, top, width, height, text, *, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run(run, size=size, bold=bold, color=color)
    return tf


def _add_bullets(slide, left, top, width, height, items, *, size=16, color=INK):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = f"•  {item}"
        _set_run(run, size=size, color=color)
    return tf


def _bar(slide, prs):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
    shape.fill.solid()
    shape.fill.fore_color.rgb = TEAL
    shape.line.fill.background()


def _footer(slide, prs, page: int, total: int = 10):
    _add_textbox(
        slide,
        Inches(0.5),
        prs.slide_height - Inches(0.45),
        Inches(8),
        Inches(0.3),
        "MindGrip AI · so101-lab · © 2026 Subhendu Datta Bhowmik",
        size=10,
        color=MUTED,
    )
    _add_textbox(
        slide,
        prs.slide_width - Inches(1.2),
        prs.slide_height - Inches(0.45),
        Inches(0.8),
        Inches(0.3),
        f"{page}/{total}",
        size=10,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def _title_block(slide, title: str, subtitle: str | None = None):
    _add_textbox(slide, Inches(0.55), Inches(0.35), Inches(12), Inches(0.55), title, size=28, bold=True, color=INK)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.95), Inches(1.4), Inches(0.06))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT_LINE
    line.line.fill.background()
    if subtitle:
        _add_textbox(slide, Inches(0.55), Inches(1.1), Inches(12), Inches(0.4), subtitle, size=14, color=MUTED)


def _maybe_picture(slide, path: Path, left, top, width, height=None):
    if path.is_file():
        if height is None:
            slide.shapes.add_picture(str(path), left, top, width=width)
        else:
            slide.shapes.add_picture(str(path), left, top, width=width, height=height)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # ── 1 Title ──
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = INK
    bg.line.fill.background()
    accent = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.22), prs.slide_height)
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()
    _maybe_picture(s, BRAND, Inches(0.7), Inches(1.4), Inches(2.4))
    _add_textbox(s, Inches(0.7), Inches(3.0), Inches(11), Inches(0.7), "MindGrip AI", size=40, bold=True, color=WHITE)
    _add_textbox(
        s,
        Inches(0.7),
        Inches(3.7),
        Inches(11),
        Inches(0.8),
        "SO-101 Pick & Place · Binary Sort Lab",
        size=24,
        color=TEAL,
    )
    _add_textbox(
        s,
        Inches(0.7),
        Inches(4.6),
        Inches(11),
        Inches(0.8),
        "Curriculum & tooling overlay on Hugging Face LeRobot — assemble → calibrate → teleop → record → train ACT → deploy",
        size=15,
        color=RGBColor(0xB8, 0xC4, 0xCE),
    )
    _add_textbox(
        s,
        Inches(0.7),
        Inches(6.5),
        Inches(11),
        Inches(0.35),
        "Subhendu Datta Bhowmik  ·  so101-lab  ·  2026",
        size=12,
        color=MUTED,
    )

    # ── 2 Problem & outcome ──
    s = prs.slides.add_slide(blank)
    _bar(s, prs)
    _title_block(s, "Problem & defined outcome", "From docs · Project Description")
    _add_bullets(
        s,
        Inches(0.55),
        Inches(1.7),
        Inches(7.2),
        Inches(4.5),
        [
            "Teach an SO-101 dual-arm kit to sort objects from demonstration — not hand-coded paths",
            "Primary task: match-box objects → match-box bin; all other objects → reject bin",
            "Bridge teleop → dataset → ACT policy → real-robot evaluation",
            "Constraints: fragile 3D-printed parts, limited kit docs, need a vision + teleop pipeline",
            "Outcome: working binary sorter with reproducible LeLab / CLI workflow + dataset + checkpoint",
        ],
        size=17,
    )
    _maybe_picture(s, HERO, Inches(8.2), Inches(1.7), Inches(4.5))
    _footer(s, prs, 2)

    # ── 3 Use case ──
    s = prs.slides.add_slide(blank)
    _bar(s, prs)
    _title_block(s, "Primary use case — binary sort", "Same pipeline later extends to colour / shape")
    cards = [
        ("Match class", "Match-box-like objects\n→ match-box bin"),
        ("Reject class", "All other objects\n→ reject bin"),
        ("Data plan", "~40–50 demos / class\n~80–100 total episodes"),
        ("Policy", "ACT (Action Chunking\nTransformer) — primary"),
    ]
    for i, (h, body) in enumerate(cards):
        left = Inches(0.55 + i * 3.15)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(2.95), Inches(3.2))
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT
        card.line.color.rgb = RGBColor(0xD0, 0xD8, 0xDE)
        top = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(1.8), Inches(2.95), Inches(0.12))
        top.fill.solid()
        top.fill.fore_color.rgb = TEAL
        top.line.fill.background()
        _add_textbox(s, left + Inches(0.2), Inches(2.15), Inches(2.55), Inches(0.5), h, size=16, bold=True, color=TEAL_DEEP)
        _add_textbox(s, left + Inches(0.2), Inches(2.8), Inches(2.55), Inches(1.8), body, size=15, color=SLATE)
    _footer(s, prs, 3)

    # ── 4 Pipeline ──
    s = prs.slides.add_slide(blank)
    _bar(s, prs)
    _title_block(s, "End-to-end pipeline", "MindGrip path on top of LeRobot")
    steps = ["Assemble", "Calibrate", "Teleop", "Record", "Train ACT", "Deploy"]
    for i, label in enumerate(steps):
        left = Inches(0.55 + i * 2.1)
        oval = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.3), Inches(1.85), Inches(1.1))
        oval.fill.solid()
        oval.fill.fore_color.rgb = TEAL if i % 2 == 0 else TEAL_DEEP
        oval.line.fill.background()
        _add_textbox(s, left, Inches(2.55), Inches(1.85), Inches(0.6), label, size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            _add_textbox(s, left + Inches(1.75), Inches(2.55), Inches(0.4), Inches(0.5), "→", size=20, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
    _add_bullets(
        s,
        Inches(0.55),
        Inches(4.0),
        Inches(12),
        Inches(2.2),
        [
            "Control host: laptop/desktop (8–16 GB) or Raspberry Pi for teleop, record, and rollout",
            "ACT training (Step 8) on cloud NVIDIA GPU / HF Jobs — not on the Air or Pi",
            "Prefer LeLab GUI after assembly; use ./so101_cli_menu.sh on Pi or for full control",
        ],
        size=16,
    )
    _footer(s, prs, 4)

    # ── 5 Hardware ──
    s = prs.slides.add_slide(blank)
    _bar(s, prs)
    _title_block(s, "Hardware — SO-101 dual arm", "Leader (teleop) + follower (policy / camera)")
    _add_bullets(
        s,
        Inches(0.55),
        Inches(1.7),
        Inches(6.5),
        Inches(4.5),
        [
            "Leader: 5 V Feetech bus — human demonstration arm (id LeoArm)",
            "Follower: 12 V Feetech bus — executes teleop then ACT policy (id LeoArm)",
            "USB camera on follower only (OpenCV front · 640×480 @ 30 FPS)",
            "Fixed scene: ~30×30 cm mat, match + reject bins, consistent camera pose",
            "Calibrate both arms before teleop; save LeoArm.json under HF_LEROBOT_HOME",
        ],
        size=16,
    )
    _maybe_picture(s, SCENE, Inches(7.4), Inches(1.6), Inches(5.3))
    _footer(s, prs, 5)

    # ── 6 Software ──
    s = prs.slides.add_slide(blank)
    _bar(s, prs)
    _title_block(s, "Software stack — so101-lab overlay", "Vanilla LeRobot stays the core")
    left_items = [
        "lerobot/ — upstream core (sibling checkout)",
        "so101-lab/ — plugins, CLI, patches, docs only",
        "MuJoCo + keyboard teleop as optional plugins",
        "Optional UX patches (TTY keys, OBS cam, viz)",
    ]
    right_items = [
        "LeLab — calibrate → teleop → record → train → deploy",
        "so101_cli_menu.sh — ports, cam, record, viz, train hooks",
        "Config: .so101_cli.env (menu C)",
        "Docs: docs-web SPA (roadmap, kinematics, CLI)",
    ]
    _add_textbox(s, Inches(0.55), Inches(1.65), Inches(5.8), Inches(0.4), "Repo layout", size=16, bold=True, color=TEAL_DEEP)
    _add_bullets(s, Inches(0.55), Inches(2.1), Inches(5.8), Inches(3.5), left_items, size=15)
    _add_textbox(s, Inches(6.9), Inches(1.65), Inches(5.8), Inches(0.4), "Day-to-day tools", size=16, bold=True, color=TEAL_DEEP)
    _add_bullets(s, Inches(6.9), Inches(2.1), Inches(5.8), Inches(3.5), right_items, size=15)
    _footer(s, prs, 6)

    # ── 7 Roadmap ──
    s = prs.slides.add_slide(blank)
    _bar(s, prs)
    _title_block(s, "Roadmap — 9 steps", "From zero hardware to autonomous sort")
    roadmap = [
        ("1", "Install & host setup"),
        ("2", "Mechanical assembly"),
        ("3", "Motor mid-pose / IDs"),
        ("4", "Calibrate both arms"),
        ("5", "Teleop verify"),
        ("6", "Scene + camera"),
        ("7", "Record demos"),
        ("8", "Train ACT (cloud)"),
        ("9", "Rollout / evaluate"),
    ]
    for i, (n, label) in enumerate(roadmap):
        row, col = divmod(i, 3)
        left = Inches(0.55 + col * 4.2)
        top = Inches(1.7 + row * 1.55)
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(3.9), Inches(1.25))
        box.fill.solid()
        box.fill.fore_color.rgb = LIGHT
        box.line.color.rgb = RGBColor(0xD0, 0xD8, 0xDE)
        badge = s.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.2), top + Inches(0.35), Inches(0.55), Inches(0.55))
        badge.fill.solid()
        badge.fill.fore_color.rgb = TEAL
        badge.line.fill.background()
        _add_textbox(s, left + Inches(0.2), top + Inches(0.42), Inches(0.55), Inches(0.4), n, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        _add_textbox(s, left + Inches(0.95), top + Inches(0.42), Inches(2.7), Inches(0.5), label, size=16, bold=True, color=INK)
    _footer(s, prs, 7)

    # ── 8 Data & ACT ──
    s = prs.slides.add_slide(blank)
    _bar(s, prs)
    _title_block(s, "Data collection & ACT training", "Steps 7–8")
    _add_bullets(
        s,
        Inches(0.55),
        Inches(1.65),
        Inches(7.0),
        Inches(4.5),
        [
            "Record with USB front camera + leader–follower teleop (CLI · 10 or LeLab)",
            "Each episode: joint state/action + images + timestamps + task string",
            "Visualize episodes (CLI · 11 / Rerun) and drop bad demos before train",
            "Train ACT on HF Jobs / cloud GPU — maps vision + state → action chunks",
            "Push dataset & checkpoint to Hugging Face Hub for reuse and eval",
        ],
        size=16,
    )
    _maybe_picture(s, CONTROL, Inches(8.0), Inches(1.7), Inches(4.7))
    _footer(s, prs, 8)

    # ── 9 Deploy & success ──
    s = prs.slides.add_slide(blank)
    _bar(s, prs)
    _title_block(s, "Deploy & success criteria", "Step 9 · Sense → ACT → Actuate")
    _add_bullets(
        s,
        Inches(0.55),
        Inches(1.7),
        Inches(12),
        Inches(4.5),
        [
            "Roll out the trained policy on the follower with the same camera pose as recording",
            "Success: correct-bin placement for match vs reject classes on held-out scenes",
            "Deliverables: reproducible workflow, dataset, ACT checkpoint, and project docs",
            "Optional sim: MuJoCo viewer / leader→sim (CLI · 16–18) — not required for the sorter",
            "Control flow: camera + joints → ACT inference → follower motor commands",
        ],
        size=17,
    )
    _footer(s, prs, 9)

    # ── 10 Resources ──
    s = prs.slides.add_slide(blank)
    _bar(s, prs)
    _title_block(s, "Resources & next steps", "Start from the docs site")
    _add_bullets(
        s,
        Inches(0.55),
        Inches(1.7),
        Inches(12),
        Inches(4.0),
        [
            "Local docs: cd docs-web && npm run dev → http://localhost:8000/",
            "Key pages: Roadmap · Project · LeLab · CLI Menu · Policy Cheatsheet · Kinematics",
            "CLI: ./so101_cli_menu.sh  ·  Configure with C  ·  Record with 10  ·  Viz with 11",
            "LeRobot docs: https://huggingface.co/docs/lerobot  ·  LeLab: /docs/lerobot/lelab",
            "Author: Subhendu Datta Bhowmik · https://subhdb.co.in",
        ],
        size=17,
    )
    _add_textbox(
        s,
        Inches(0.55),
        Inches(5.8),
        Inches(12),
        Inches(0.5),
        "Finish the binary match-box sorter first — then reuse the same pipeline for colour / shape scenarios.",
        size=15,
        bold=True,
        color=TEAL_DEEP,
    )
    _footer(s, prs, 10)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path} ({path.stat().st_size // 1024} KB, {10} slides)")
